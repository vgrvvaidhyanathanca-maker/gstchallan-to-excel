"""
gst_challan_extract.py
----------------------
Pulls the key fields out of GST challans (PDF) and writes them to one Excel
sheet, one row per CPIN. Handles both layouts the GST portal produces:

  * Form GST PMT-06  (the unpaid challan you generate before paying)
  * GST PAYMENT RECEIPT (the paid receipt, with CIN / bank / BRN)

Usage (from a terminal):
    py gst_challan_extract.py                        -> scans the folder the script is in
    py gst_challan_extract.py "D:\\Clients\\Challans"  -> scans that folder (and sub-folders)
    py gst_challan_extract.py "D:\\Challans" out.xlsx  -> custom output file

Re-running is safe: the sheet is rebuilt from existing rows + new PDFs, keyed
on CPIN. If a CPIN appears both as a PMT-06 form and as a Payment Receipt, the
receipt wins (it has the payment details). Non-GST PDFs are ignored.

Needs: pymupdf, openpyxl   (py -m pip install pymupdf openpyxl)
"""

import re
import sys
from datetime import datetime
from pathlib import Path

import fitz  # pymupdf
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Output layout
# ----------------------------------------------------------------------------
TAX_HEADS = ["CGST", "IGST", "CESS", "SGST"]          # SGST also covers UTGST
TAX_COLS = ["Tax", "Interest", "Penalty", "Fee", "Others", "Total"]

COLUMNS = (
    ["File", "Type", "Status", "CPIN", "Challan Date", "Expiry Date", "GSTIN",
     "Legal Name", "State", "Reason", "Financial Year", "Period"]
    + [f"{h} {c}" for h in TAX_HEADS for c in TAX_COLS]
    + ["Total Tax", "Total Interest", "Total Penalty", "Total Fee",
       "Total Others", "Grand Total", "Mode of Payment",
       "Bank", "BRN / UTR", "CIN", "Payment Date"]
)
AMOUNT_COLS = {c for c in COLUMNS
               if any(c.endswith(" " + t) for t in TAX_COLS) or c == "Grand Total"}
DATE_COLS = {"Challan Date", "Expiry Date", "Payment Date"}

MODE_LABELS = {          # first word of each option as printed on the PMT-06
    "E-Payment": "E-Payment",
    "Over": "Over the Counter (OTC)",
    "NEFT": "NEFT / RTGS",
}


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def to_number(s):
    """'20,108' -> 20108 ; '-' or '' -> 0"""
    s = (s or "").replace(",", "").strip()
    if s in ("", "-"):
        return 0
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return s


def to_date(s):
    """'17/08/2026' -> datetime ; anything else -> unchanged / ''"""
    s = (s or "").strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return s


def words_on_row(words, y, tol=3):
    """All words whose vertical centre is within `tol` points of y."""
    return sorted(
        (w for w in words if abs((w[1] + w[3]) / 2 - y) <= tol),
        key=lambda w: w[0],
    )


def text_between(text, label, stops):
    """Text after `label` up to the first of the `stops` labels that follows.
    Whitespace/newlines are collapsed, so wrapped values come out on one line."""
    i = text.find(label)
    if i < 0:
        return ""
    start = i + len(label)
    end = len(text)
    for s in stops:
        j = text.find(s, start)
        if 0 <= j < end:
            end = j
    return re.sub(r"\s+", " ", text[start:end]).strip()


def first_line_after(text, label):
    m = re.search(re.escape(label) + r"[ \t]*([^\n]*)", text)
    return m.group(1).strip() if m else ""


# ----------------------------------------------------------------------------
# Core parser
# ----------------------------------------------------------------------------
def parse_challan(pdf_path: Path) -> dict | None:
    doc = fitz.open(pdf_path)
    page = doc[0]
    text = page.get_text()

    if "PMT" in text and "CPIN" in text and "Payment Challan" in text:
        kind = "PMT-06 Form"
    elif "PAYMENT RECEIPT" in text and "CPIN" in text and "GSTIN" in text:
        kind = "Payment Receipt"
    else:
        return None                       # not a GST challan

    words = page.get_text("words")        # (x0, y0, x1, y1, text, block, line, wordno)
    row = {c: "" for c in COLUMNS}
    row["File"] = pdf_path.name
    row["Type"] = kind
    row["Status"] = "Paid" if kind == "Payment Receipt" else "Unpaid / Generated"

    # ---- header block ----------------------------------------------------
    m = re.search(r"CPIN:\s*(\d{14})", text)
    row["CPIN"] = m.group(1) if m else ""
    m = re.search(r"GSTIN:\s*([0-9A-Z]{15})", text)
    row["GSTIN"] = m.group(1) if m else ""

    if kind == "PMT-06 Form":
        m = re.search(r"Challan Generated on\s*:\s*(\d{2}/\d{2}/\d{4})", text)
        row["Challan Date"] = to_date(m.group(1) if m else "")
        m = re.search(r"Expiry Date\s*:\s*(\d{2}/\d{2}/\d{4})", text)
        row["Expiry Date"] = to_date(m.group(1) if m else "")
        row["Legal Name"] = text_between(text, "Name(Legal):", ["Address"])
    else:
        m = re.search(r"Deposit Date\s*:\s*(\d{2}/\d{2}/\d{4})", text)
        row["Challan Date"] = to_date(m.group(1) if m else "")
        row["Payment Date"] = row["Challan Date"]
        row["Legal Name"] = text_between(text, "Name:", ["Address"])
        row["CIN"] = first_line_after(text, "CIN:")
        row["Bank"] = text_between(text, "Name of Bank:", ["BRN:", "Details of Taxpayer"])
        row["BRN / UTR"] = first_line_after(text, "BRN:")
        row["Mode of Payment"] = first_line_after(text, "Mode of Payment:")

    # Reason may wrap on to several lines and FY / Period may be absent
    row["Reason"] = text_between(text, "Reason:", ["Financial Year", "Period:", "Details of Deposit"])
    row["Financial Year"] = text_between(text, "Financial Year:", ["Period:", "Details of Deposit"])
    row["Period"] = text_between(text, "Period:", ["Details of Deposit"])

    # ---- deposit table (position based) ----------------------------------
    # column x-positions come from the "Tax Interest Penalty Fee Others Total" header row
    hdr_x = {}
    anchor = next((w for w in words if w[4] == "Interest"), None)
    if anchor:
        for w in words_on_row(words, (anchor[1] + anchor[3]) / 2):
            if w[4] in TAX_COLS:
                hdr_x.setdefault(w[4], (w[0] + w[2]) / 2)

    def nearest_col(x):
        return min(hdr_x, key=lambda c: abs(hdr_x[c] - x)) if hdr_x else None

    state = ""
    for w in words:
        m = re.match(r"^(CGST|IGST|CESS|SGST|UTGST)\(\d{4}\)$", w[4])
        if not m:
            continue
        head = "SGST" if m.group(1) == "UTGST" else m.group(1)
        y = (w[1] + w[3]) / 2
        for v in words_on_row(words, y):
            if v[0] <= w[2]:                          # skip the label and anything left of it
                continue
            col = nearest_col((v[0] + v[2]) / 2)
            if col:
                row[f"{head} {col}"] = to_number(v[4])
        # the state name sits immediately left of the SGST/UTGST label
        if head == "SGST":
            left = [v[4] for v in words_on_row(words, y, tol=8) if v[2] <= w[0]]
            state = " ".join(left)
    row["State"] = state

    for h in TAX_HEADS:
        for c in TAX_COLS:
            if row[f"{h} {c}"] == "":
                row[f"{h} {c}"] = 0
    for c in TAX_COLS[:-1]:
        row[f"Total {c}"] = sum(row[f"{h} {c}"] for h in TAX_HEADS
                                if isinstance(row[f"{h} {c}"], (int, float)))
    m = re.search(r"Total Amount\s*\n\s*([\d,]+)", text)
    row["Grand Total"] = to_number(m.group(1)) if m else sum(
        row[f"{h} Total"] for h in TAX_HEADS if isinstance(row[f"{h} Total"], (int, float)))

    # ---- mode of payment on the PMT-06 (tick-box image, not text) ---------
    # An empty box has ink only on its border, so we look at the centre of the
    # box beside each option: the one with ink in the middle carries the tick.
    if kind == "PMT-06 Form":
        best, best_ink = "", 0
        for w in words:
            if w[4] in MODE_LABELS:
                box = fitz.Rect(w[0] - 20, w[1] - 1, w[0] - 1, w[3] + 1)
                inner = fitz.Rect(box.x0 + box.width * 0.3, box.y0 + box.height * 0.3,
                                  box.x1 - box.width * 0.3, box.y1 - box.height * 0.3)
                pix = page.get_pixmap(clip=inner, dpi=300, colorspace=fitz.csGRAY)
                ink = sum(1 for b in pix.samples if b < 160)
                if ink > best_ink:
                    best, best_ink = MODE_LABELS[w[4]], ink
        row["Mode of Payment"] = best

        # "Paid Challan Information" block is filled only on a paid PMT-06
        paid_hdr = next((w for w in words if w[4] == "Paid" and
                         any(v[4] == "Challan" for v in words_on_row(words, (w[1] + w[3]) / 2))), None)
        if paid_hdr:
            lower = [w for w in words if w[1] > paid_hdr[3]]

            def value_of(first_word):
                for w in lower:
                    if w[4] == first_word:
                        y = (w[1] + w[3]) / 2
                        return " ".join(v[4] for v in words_on_row(lower, y) if v[0] > 250)
                return ""

            row["Bank"] = value_of("Name")
            brn = value_of("Bank")
            row["BRN / UTR"] = "" if "Ack" in brn else brn
            row["CIN"] = value_of("CIN")
            row["Payment Date"] = to_date(value_of("Payment"))
            if row["CIN"]:
                row["Status"] = "Paid"

    return row


# ----------------------------------------------------------------------------
# Excel writer
# ----------------------------------------------------------------------------
def load_existing(path: Path) -> dict:
    """Existing rows keyed by CPIN (so a re-run rebuilds rather than duplicates)."""
    if not path.exists():
        return {}
    ws = load_workbook(path).active
    hdr = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        if d.get("CPIN"):
            out[str(d["CPIN"])] = {c: d.get(c, "") for c in COLUMNS}
    return out


def merge(existing: dict, new_rows: list) -> dict:
    """Receipt beats form; otherwise the newer parse replaces the old one."""
    merged = dict(existing)
    for r in new_rows:
        cur = merged.get(r["CPIN"])
        if cur and cur.get("Type") == "Payment Receipt" and r["Type"] != "Payment Receipt":
            continue
        merged[r["CPIN"]] = r
    return merged


def write_excel(rows: dict, out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "GST Challans"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "E2"

    def sort_key(r):
        d = r.get("Challan Date")
        return d if isinstance(d, datetime) else datetime.min

    for r in sorted(rows.values(), key=sort_key):
        ws.append([r.get(c, "") for c in COLUMNS])

    for i, col in enumerate(COLUMNS, 1):
        letter = get_column_letter(i)
        if col in AMOUNT_COLS:
            width = 12
            for r in range(2, ws.max_row + 1):
                ws.cell(r, i).number_format = "#,##0"
        elif col in DATE_COLS:
            width = 12
            for r in range(2, ws.max_row + 1):
                ws.cell(r, i).number_format = "dd/mm/yyyy"
        else:
            width = max(len(col), *(len(str(ws.cell(r, i).value or ""))
                                    for r in range(2, ws.max_row + 1)))
            width = min(max(width + 2, 8), 45)
        ws.column_dimensions[letter].width = width
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)


# ----------------------------------------------------------------------------
def main():
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else folder / "GST_Challans.xlsx"

    pdfs = sorted(folder.rglob("*.pdf"))
    if not pdfs:
        print(f"No PDF files found in {folder}")
        return

    new_rows, skipped = [], 0
    for pdf in pdfs:
        try:
            r = parse_challan(pdf)
        except Exception as e:  # keep going on a bad file
            print(f"  ERROR {pdf.name}: {e}")
            continue
        if r is None:
            skipped += 1
            continue
        print(f"  {pdf.name}: {r['Type']} | {r['Legal Name']} | {r['Financial Year']} {r['Period']} | "
              f"Rs {r['Grand Total']} | {r['Mode of Payment']}")
        new_rows.append(r)

    existing = load_existing(out)
    merged = merge(existing, new_rows)
    write_excel(merged, out)
    print(f"\n{len(new_rows)} GST challan PDF(s) read, {skipped} other PDF(s) ignored.")
    print(f"{len(merged)} unique challan(s) (by CPIN) written -> {out}")


if __name__ == "__main__":
    main()
